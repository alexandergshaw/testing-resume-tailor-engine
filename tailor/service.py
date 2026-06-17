"""The one entry point every consumer shares: web UI, API, tests.

Stateless and in-memory: takes posting text + docx bytes, returns proposals
or a tailored docx. Optional per-request profile/library override the
bundled data files, which makes the API multi-tenant-friendly.
"""
import io
from dataclasses import dataclass

import docx

from .clients.base import DownstreamError
from .docxio.filler import fill
from .docxio.generator_render import DOCX_MIME, prepare_for_generator
from .docxio.scanner import scan
from .extraction.extractor import Keyword, extract_keywords
from .library.store import LibraryEntry, load_library, load_profile, slugify
from .mapping.strategies import propose_all
from .paths import DEFAULT_COVER_LETTER_TEMPLATE


class ServiceError(Exception):
    """Base for caller errors; `status` maps to the HTTP response code."""
    status = 400


class InvalidInputError(ServiceError):
    status = 400


class InvalidDocxError(ServiceError):
    status = 400


class NoPlaceholdersError(ServiceError):
    status = 422


@dataclass(frozen=True)
class SlotProposal:
    key: str            # "NAME::occurrence" — stable id used in values overrides
    name: str
    occurrence: int
    raw: str
    context: str
    single_brace: bool
    strategy: str
    value: str
    note: str
    candidates: tuple[str, ...]

    def to_dict(self) -> dict:
        return {
            "key": self.key, "name": self.name, "occurrence": self.occurrence,
            "raw": self.raw, "context": self.context,
            "single_brace": self.single_brace, "strategy": self.strategy,
            "value": self.value, "note": self.note,
            "candidates": list(self.candidates),
        }


def parse_library_entries(raw_entries: list) -> list[LibraryEntry]:
    """Build LibraryEntry objects from request JSON; ids are optional."""
    entries = []
    taken: set[str] = set()
    for index, raw in enumerate(raw_entries):
        if not isinstance(raw, dict) or not raw.get("text"):
            raise InvalidInputError(f"library entry {index} needs at least 'text'")
        base = str(raw.get("id") or slugify(str(raw["text"])))
        entry_id = base
        counter = 2
        while entry_id in taken:
            entry_id = f"{base}-{counter}"
            counter += 1
        taken.add(entry_id)
        entries.append(LibraryEntry(
            id=entry_id,
            slots=tuple(str(s) for s in raw.get("slots", [])),
            text=str(raw["text"]),
            tags=tuple(str(t) for t in raw.get("tags", [])),
            metric=str(raw.get("metric", "")),
        ))
    return entries


def _open_document(docx_bytes: bytes):
    try:
        return docx.Document(io.BytesIO(docx_bytes))
    except Exception as exc:
        raise InvalidDocxError(f"template could not be opened as a .docx: {exc}") from exc


def _resolve_data(profile, library):
    if profile is not None and not isinstance(profile, dict):
        raise InvalidInputError("'profile' must be an object of placeholder -> value")
    if library is not None and not isinstance(library, list):
        raise InvalidInputError("'library' must be a list of entries")
    resolved_profile = ({str(k): str(v) for k, v in profile.items()}
                        if profile is not None else load_profile())
    resolved_library = (parse_library_entries(library)
                        if library is not None else load_library())
    return resolved_profile, resolved_library


def propose_for(posting: str, docx_bytes: bytes, profile: dict | None = None,
                library: list | None = None,
                keywords: list[Keyword] | None = None
                ) -> tuple[list[SlotProposal], list[Keyword]]:
    """`keywords`, when provided, replaces local extraction — this is how the
    composed workflow injects Parser-derived keywords. When None (legacy /
    default), keywords are extracted locally, identical to the original engine."""
    if not posting or not posting.strip():
        raise InvalidInputError("'posting' must not be empty")
    resolved_profile, resolved_library = _resolve_data(profile, library)

    document = _open_document(docx_bytes)
    placeholders = scan(document)
    if not placeholders:
        raise NoPlaceholdersError("no {{placeholders}} found in the template")

    if keywords is None:
        keywords = extract_keywords(posting)
    proposals = propose_all(placeholders, keywords, resolved_profile, resolved_library)

    slots = [
        SlotProposal(
            key=ph.key, name=ph.name, occurrence=ph.occurrence, raw=ph.raw,
            context=ph.context, single_brace=ph.single_brace,
            strategy=proposals[ph.key].strategy.value,
            value=proposals[ph.key].value,
            note=proposals[ph.key].note,
            candidates=proposals[ph.key].candidates,
        )
        for ph in placeholders
    ]
    return slots, keywords


def _render_local(docx_bytes: bytes, fill_values: dict[tuple[str, int], str]) -> bytes:
    document = _open_document(docx_bytes)
    fill(document, fill_values)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _render_via_generator(docx_bytes: bytes, fill_values: dict[tuple[str, int], str],
                          generator_client) -> bytes:
    import json
    template_bytes, content = prepare_for_generator(docx_bytes, fill_values)
    return generator_client.generate(
        "docx", json.dumps(content),
        template_file=("resume.docx", template_bytes, DOCX_MIME), strict=False)


def tailor_document(posting: str, docx_bytes: bytes, profile: dict | None = None,
                    library: list | None = None, values: dict | None = None,
                    field_values: dict | None = None,
                    keywords: list[Keyword] | None = None,
                    generator_client=None) -> tuple[bytes, dict]:
    """Fill the template. `values` maps slot key ("NAME::occ") -> final text;
    `field_values` maps a placeholder NAME -> text applied to ALL its
    occurrences (handy for per-applicant fields like TARGET_ORGANIZATION that
    repeat). Precedence per slot: values > field_values > proposal. An empty
    result leaves the {{placeholder}} in the document and is reported unfilled.
    `keywords` injects a precomputed keyword set (composed workflow).
    `generator_client`, when given, renders the final docx via the Document
    Generator API (composed workflow), falling back to the local filler if it
    is unavailable."""
    if values is not None and not isinstance(values, dict):
        raise InvalidInputError("'values' must be an object of slot key -> text")
    if field_values is not None and not isinstance(field_values, dict):
        raise InvalidInputError("'field_values' must be an object of name -> text")
    slots, keywords = propose_for(posting, docx_bytes, profile, library, keywords=keywords)

    fill_values: dict[tuple[str, int], str] = {}
    report_slots = []
    for slot in slots:
        override = None if values is None else values.get(slot.key)
        field_override = None if field_values is None else field_values.get(slot.name)
        if override is not None:
            final = str(override).replace("\r\n", "\n").strip()
            source = "overridden"
        elif field_override is not None:
            final = str(field_override).replace("\r\n", "\n").strip()
            source = "field"
        else:
            final = slot.value
            source = "proposed"
        if final:
            fill_values[(slot.name, slot.occurrence)] = final
        else:
            source = "unfilled"
        report_slots.append({**slot.to_dict(), "final_value": final, "source": source})

    report = {
        "slots": report_slots,
        "unfilled": [s["key"] for s in report_slots if s["source"] == "unfilled"],
        "keywords": keywords_payload(keywords),
    }

    rendered, renderer = None, "local"
    if generator_client is not None:
        try:
            rendered = _render_via_generator(docx_bytes, fill_values, generator_client)
            renderer = "generator"
        except DownstreamError as exc:
            report.setdefault("warnings", []).append(
                f"generator unavailable ({exc}); rendered locally")
    if rendered is None:
        rendered = _render_local(docx_bytes, fill_values)
    report.setdefault("meta", {})["renderer"] = renderer
    return rendered, report


def tailor_cover_letter(posting: str, docx_bytes: bytes | None = None,
                        target_role: str | None = None,
                        target_organization: str | None = None,
                        profile: dict | None = None, library: list | None = None,
                        values: dict | None = None,
                        keywords: list[Keyword] | None = None,
                        extra_field_values: dict | None = None,
                        generator_client=None) -> tuple[bytes, dict]:
    """Cover-letter convenience over tailor_document: falls back to the bundled
    template when none is supplied, expands target_role/target_organization
    across every {{TARGET_ROLE}}/{{TARGET_ORGANIZATION}} occurrence, and applies
    `extra_field_values` (e.g. research-composed ORGANIZATION_CONTEXT/ROLE_FOCUS)
    to all occurrences of those names."""
    if docx_bytes is None:
        docx_bytes = DEFAULT_COVER_LETTER_TEMPLATE.read_bytes()
    field_values = {}
    if target_role is not None:
        field_values["TARGET_ROLE"] = str(target_role)
    if target_organization is not None:
        field_values["TARGET_ORGANIZATION"] = str(target_organization)
    if extra_field_values:
        field_values.update({k: str(v) for k, v in extra_field_values.items() if v})
    return tailor_document(posting, docx_bytes, profile=profile, library=library,
                           values=values, field_values=field_values or None,
                           keywords=keywords, generator_client=generator_client)


def keywords_payload(keywords: list[Keyword]) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = {}
    for kw in keywords:
        grouped.setdefault(kw.category, []).append(
            {"canonical": kw.canonical, "score": kw.score, "count": kw.count})
    return grouped
