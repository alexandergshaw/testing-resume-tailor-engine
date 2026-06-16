"""Adapt an occurrence-indexed template for the generic (Jinja/docxtpl) Document
Generator API.

The generator renders one value per variable name, so it can't fill five
identical {{SKILLS_LINE}} placeholders with five different values. We rewrite
each occurrence to a unique Jinja variable ({{SKILLS_LINE__0}}, ...) and build a
matching flat content dict. Unfilled slots keep their original token text so
they stay visible in the output, exactly like the local filler.
"""
import io

import docx

from .scanner import scan
from .filler import fill

DOCX_MIME = ("application/vnd.openxmlformats-officedocument"
             ".wordprocessingml.document")


def _var_name(name: str, occurrence: int) -> str:
    return f"{name}__{occurrence}"


def prepare_for_generator(docx_bytes: bytes,
                          fill_values: dict[tuple[str, int], str]
                          ) -> tuple[bytes, dict[str, str]]:
    """Returns (transformed_template_bytes, content). The template has unique
    Jinja variables; content maps each variable to its value (or the original
    placeholder token when unfilled)."""
    document = docx.Document(io.BytesIO(docx_bytes))
    placeholders = scan(document)

    rename: dict[tuple[str, int], str] = {}
    content: dict[str, str] = {}
    for ph in placeholders:
        var = _var_name(ph.name, ph.occurrence)
        rename[(ph.name, ph.occurrence)] = "{{" + var + "}}"
        value = fill_values.get((ph.name, ph.occurrence))
        content[var] = value if value else ph.raw  # unfilled -> keep token visible

    fill(document, rename)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), content
