from pathlib import Path

import docx
from flask import (Blueprint, abort, flash, redirect, render_template, request,
                   send_file, url_for)

from ..docxio.filler import fill
from ..docxio.scanner import scan
from ..extraction.extractor import extract_keywords, keywords_by_category
from ..library.store import load_library, load_profile
from ..mapping.strategies import propose_all
from ..paths import UPLOADS_DIR
from . import sessions
from .sessions import ParseSession

bp = Blueprint("web", __name__)


@bp.get("/")
def index():
    return render_template("index.html")


@bp.post("/parse")
def parse():
    posting = request.form.get("posting", "").strip()
    upload = request.files.get("template")

    if not posting:
        flash("Paste a job posting first.")
        return redirect(url_for("web.index"))
    if upload is None or not upload.filename:
        flash("Upload your template .docx.")
        return redirect(url_for("web.index"))
    if not upload.filename.lower().endswith(".docx"):
        flash("The template must be a .docx file.")
        return redirect(url_for("web.index"))

    session = ParseSession(docx_path=Path(), original_filename=upload.filename,
                           placeholders=[], proposals={}, keywords=[])
    token = sessions.create(session)

    UPLOADS_DIR.mkdir(exist_ok=True)
    docx_path = UPLOADS_DIR / f"{token}.docx"
    upload.save(docx_path)

    try:
        document = docx.Document(docx_path)
    except Exception:
        flash("That file could not be opened as a Word document.")
        return redirect(url_for("web.index"))

    placeholders = scan(document)
    if not placeholders:
        flash("No {{placeholders}} were found in that document.")
        return redirect(url_for("web.index"))

    keywords = extract_keywords(posting)
    proposals = propose_all(placeholders, keywords, load_profile(), load_library())

    session.docx_path = docx_path
    session.placeholders = placeholders
    session.proposals = proposals
    session.keywords = keywords
    return redirect(url_for("web.review", token=token))


@bp.get("/review/<token>")
def review(token):
    session = sessions.get(token)
    if session is None or not session.placeholders:
        flash("That session expired (the app was restarted). Start again.")
        return redirect(url_for("web.index"))
    rows = [(ph, session.proposals[ph.key]) for ph in session.placeholders]
    manual_count = sum(1 for _, prop in rows if not prop.value)
    return render_template(
        "review.html",
        token=token,
        rows=rows,
        manual_count=manual_count,
        filename=session.original_filename,
        keyword_groups=keywords_by_category(session.keywords),
    )


@bp.post("/generate/<token>")
def generate(token):
    session = sessions.get(token)
    if session is None or not session.placeholders:
        abort(404)
    values = {}
    for ph in session.placeholders:
        raw = request.form.get(f"v_{ph.key}")
        if raw is None:
            continue
        value = raw.replace("\r\n", "\n").strip()
        if value:
            # Blank fields keep their {{placeholder}} so you can spot them in Word.
            values[(ph.name, ph.occurrence)] = value

    document = docx.Document(session.docx_path)
    fill(document, values)
    out_path = UPLOADS_DIR / f"{token}_filled.docx"
    document.save(out_path)
    return send_file(out_path, as_attachment=True, download_name="Tailored Resume.docx")
