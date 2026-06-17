"""The bundled standardized resume template + resume endpoints' default-template
fallback."""
import io
import re

import docx
import pytest

from app import create_app
from tailor.docxio.scanner import scan
from tailor.paths import DEFAULT_RESUME_TEMPLATE
from tailor.service import propose_for, tailor_document
from tests.test_service import POSTING

RAW_RE = re.compile(r"\{\{[A-Z0-9_]+\}\}")


def test_bundled_template_exists_and_is_consistent():
    document = docx.Document(DEFAULT_RESUME_TEMPLATE)
    placeholders = scan(document)
    assert placeholders, "bundled template should contain placeholders"
    # Every token is {{UPPER_SNAKE_CASE}} — no Title Case, spaces, or single braces.
    for ph in placeholders:
        assert RAW_RE.fullmatch(ph.raw), f"non-standard placeholder: {ph.raw!r}"
        assert not ph.single_brace


def test_standardized_names_present_old_names_absent():
    names = {ph.name for ph in scan(docx.Document(DEFAULT_RESUME_TEMPLATE))}
    assert {"SKILLS_LINE", "COURSE_TOPICS_3", "COURSE_TOPICS_4",
            "TOP_RANK", "AREA_OF_EMPHASIS", "ROLE_SPECIFIC_EXPERTISE"} <= names
    assert "2_LINES_OF_COMMA_SEPARATED_SKILLS" not in names
    assert not any(n.startswith("LIST_OF_") for n in names)


def test_bundled_template_fills_without_leftover_placeholders():
    slots, _ = propose_for(POSTING, DEFAULT_RESUME_TEMPLATE.read_bytes())
    values = {s.key: (s.value or "X") for s in slots}
    docx_bytes, report = tailor_document(
        POSTING, DEFAULT_RESUME_TEMPLATE.read_bytes(), values=values)
    result = docx.Document(io.BytesIO(docx_bytes))
    assert scan(result) == []
    assert "Alex Shaw" in "\n".join(p.text for p in result.paragraphs)


@pytest.fixture()
def client():
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def test_proposals_falls_back_to_bundled_template(client):
    # No template part -> bundled standardized template is used.
    res = client.post("/api/v1/proposals", data={"posting": POSTING},
                      content_type="multipart/form-data")
    assert res.status_code == 200
    names = {s["name"] for s in res.get_json()["slots"]}
    assert "SKILLS_LINE" in names


def test_tailor_falls_back_to_bundled_template(client):
    res = client.post("/api/v1/resume", data={"posting": POSTING},
                      content_type="multipart/form-data")
    assert res.status_code == 200
    assert res.headers["Content-Disposition"].endswith('"Tailored Resume.docx"')
    assert res.headers["X-Unfilled-Count"].isdigit()
    # A real docx came back (zip magic bytes) with posting keywords filled in.
    assert res.data[:2] == b"PK"
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(res.data)).paragraphs)
    assert "Kubernetes" in text


def test_uploaded_template_still_overrides_default(client):
    from tests.test_service import make_docx
    template = make_docx("Only {{RANK}} here")
    res = client.post("/api/v1/proposals", data={
        "posting": POSTING, "template": (io.BytesIO(template), "t.docx")},
        content_type="multipart/form-data")
    names = {s["name"] for s in res.get_json()["slots"]}
    assert names == {"RANK"}
