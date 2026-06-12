"""Full flow through the public API with the real template: proposals ->
caller edits -> tailor -> verify the returned docx.

Uses the real template from Downloads when present; skipped otherwise so the
suite still passes on other machines.
"""
import io
import json
import re
from pathlib import Path

import docx
import pytest

from app import create_app
from tailor.docxio.scanner import scan

REAL_TEMPLATE = Path(r"C:\Users\alexa\Downloads\Template Resume.docx")
POSTING = (Path(__file__).parent / "fixtures" / "sample_posting.txt").read_text(encoding="utf-8")

pytestmark = pytest.mark.skipif(not REAL_TEMPLATE.exists(),
                                reason="real template not present on this machine")


@pytest.fixture()
def client():
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def _multipart(**extra):
    data = {"posting": POSTING,
            "template": (io.BytesIO(REAL_TEMPLATE.read_bytes()), "Template Resume.docx")}
    data.update(extra)
    return {"data": data, "content_type": "multipart/form-data"}


def test_full_flow_with_real_template(client):
    proposals = client.post("/api/v1/proposals", **_multipart())
    assert proposals.status_code == 200
    body = proposals.get_json()

    slots = body["slots"]
    skills = [s for s in slots if s["name"] == "2_LINES_OF_COMMA_SEPARATED_SKILLS"]
    assert [s["occurrence"] for s in skills] == [0, 1, 2, 3, 4]
    assert sum(1 for s in slots if s["single_brace"]) >= 2
    assert "technology" in body["keywords"]

    # Caller accepts proposals, supplying stand-ins for unfilled slots.
    values = {s["key"]: s["value"] or "FILLED-BY-TEST" for s in slots}

    tailored = client.post("/api/v1/tailor", **_multipart(values=json.dumps(values)))
    assert tailored.status_code == 200
    assert tailored.headers["Content-Disposition"].endswith('"Tailored Resume.docx"')
    assert tailored.headers["X-Unfilled-Count"] == "0"

    result = docx.Document(io.BytesIO(tailored.data))
    assert scan(result) == [], "no placeholders may survive generation"
    text = "\n".join(p.text for p in result.paragraphs)
    assert not re.search(r"\{\{", text)
    assert "Alex Shaw" in text  # original content intact
