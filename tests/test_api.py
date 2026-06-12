"""Contract tests for /api/v1 — the same surface external apps depend on."""
import base64
import io
import json

import docx
import pytest

from app import create_app
from tests.test_service import POSTING, TEMPLATE


@pytest.fixture()
def client():
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def multipart(**extra):
    data = {"posting": POSTING,
            "template": (io.BytesIO(TEMPLATE), "template.docx")}
    data.update(extra)
    return {"data": data, "content_type": "multipart/form-data"}


def test_health(client):
    body = client.get("/api/v1/health").get_json()
    assert body["status"] == "ok"
    assert body["readonly"] is False
    assert body["bank_entries"] > 0


def test_proposals_multipart(client):
    res = client.post("/api/v1/proposals", **multipart())
    assert res.status_code == 200
    body = res.get_json()
    keys = [slot["key"] for slot in body["slots"]]
    assert "JOB_RELEVANT_TECHNOLOGIES::0" in keys
    assert "tool_platform" in body["keywords"]
    assert body["engine_version"]


def test_proposals_json_base64(client):
    res = client.post("/api/v1/proposals", json={
        "posting": POSTING,
        "template_b64": base64.b64encode(TEMPLATE).decode("ascii"),
    })
    assert res.status_code == 200
    assert res.get_json()["slots"]


def test_tailor_binary_default(client):
    res = client.post("/api/v1/tailor", **multipart(
        values=json.dumps({"MEASURABLE_IMPACT::0": "a 9x speedup"})))
    assert res.status_code == 200
    assert "wordprocessingml" in res.mimetype
    assert res.headers["X-Engine-Version"]
    document = docx.Document(io.BytesIO(res.data))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "a 9x speedup" in text


def test_tailor_json_accept(client):
    res = client.post("/api/v1/tailor", headers={"Accept": "application/json"},
                      **multipart())
    assert res.status_code == 200
    body = res.get_json()
    assert base64.b64decode(body["docx_b64"])[:2] == b"PK"  # a zip = docx
    assert "unfilled" in body["report"]


def test_proposals_then_tailor_round_trip_matches_direct(client):
    proposals = client.post("/api/v1/proposals", **multipart()).get_json()
    values = {slot["key"]: slot["value"] for slot in proposals["slots"]}
    via_review = client.post("/api/v1/tailor", **multipart(values=json.dumps(values)))
    direct = client.post("/api/v1/tailor", **multipart())
    text = lambda res: "\n".join(  # noqa: E731
        p.text for p in docx.Document(io.BytesIO(res.data)).paragraphs)
    assert text(via_review) == text(direct)


def test_error_responses(client):
    no_posting = client.post("/api/v1/proposals", data={
        "posting": "  ", "template": (io.BytesIO(TEMPLATE), "t.docx")},
        content_type="multipart/form-data")
    assert no_posting.status_code == 400
    assert no_posting.get_json()["error"] == "InvalidInputError"

    bad_docx = client.post("/api/v1/proposals", data={
        "posting": POSTING, "template": (io.BytesIO(b"nope"), "t.docx")},
        content_type="multipart/form-data")
    assert bad_docx.status_code == 400

    import tests.test_service as ts
    empty = client.post("/api/v1/proposals", data={
        "posting": POSTING,
        "template": (io.BytesIO(ts.make_docx("no slots here")), "t.docx")},
        content_type="multipart/form-data")
    assert empty.status_code == 422

    bad_json_field = client.post("/api/v1/tailor", **multipart(values="{not json"))
    assert bad_json_field.status_code == 400

    wrong_type = client.post("/api/v1/proposals", data="plain text",
                             content_type="text/plain")
    assert wrong_type.status_code == 400


def test_auth_enforced_when_keys_set(client, monkeypatch):
    monkeypatch.setenv("API_KEYS", "secret-1, secret-2")
    denied = client.get("/api/v1/health")
    assert denied.status_code == 401
    allowed = client.get("/api/v1/health", headers={"X-API-Key": "secret-2"})
    assert allowed.status_code == 200


def test_readonly_blocks_bank_writes_and_remember(client, monkeypatch):
    monkeypatch.setenv("RESUME_TAILOR_READONLY", "1")
    assert client.get("/api/v1/health").get_json()["readonly"] is True

    blocked = client.post("/api/v1/bank/entries",
                          json={"text": "x", "slot": "ACTION"})
    assert blocked.status_code == 403

    # remember is silently ignored: tailor succeeds, nothing persisted.
    res = client.post("/api/v1/tailor", headers={"Accept": "application/json"},
                      **multipart(values=json.dumps({"RANK::0": "Lead"}),
                                  remember=json.dumps(["RANK::0"])))
    assert res.status_code == 200
    assert "remembered" not in res.get_json()["report"]


def test_cors_headers_present(client):
    res = client.post("/api/v1/proposals", **multipart())
    assert res.headers["Access-Control-Allow-Origin"] == "*"


def test_per_request_library_drives_candidates(client):
    library = [{"text": "Spearheaded", "slots": ["RANK"], "tags": []}]
    res = client.post("/api/v1/proposals", **multipart(library=json.dumps(library)))
    slots = {slot["key"]: slot for slot in res.get_json()["slots"]}
    assert slots["RANK::0"]["candidates"] == ["Spearheaded"]
