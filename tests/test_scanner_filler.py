import docx

from tailor.docxio.filler import fill
from tailor.docxio.scanner import scan
from tailor.mapping.registry import normalize_name


def build_fixture_doc():
    doc = docx.Document()

    doc.add_paragraph("Hello {{RANK}} engineer")

    # Placeholder deliberately split across runs the way Word does it,
    # with formatting on the opening run.
    p = doc.add_paragraph()
    r1 = p.add_run("Experienced with {{")
    r1.bold = True
    p.add_run("JOB_RELEVANT")
    p.add_run("_TECHNOLOGIES")
    p.add_run("}} and more")

    # Same placeholder twice in one paragraph plus single-brace tokens.
    doc.add_paragraph("({Area of Emphasis}, {Area of Emphasis})")

    doc.add_paragraph("{{2 lines of comma separated skills}}")
    doc.add_paragraph("{{2 lines of comma separated skills}}")

    # Inside a table: occurrence indices must keep counting.
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].paragraphs[0].add_run("Cell rank: {{RANK}}")

    return doc


def test_normalize_name():
    assert normalize_name("{{Role-Specific Expertise }}") == "ROLE_SPECIFIC_EXPERTISE"
    assert normalize_name("{Area of Emphasis}") == "AREA_OF_EMPHASIS"
    assert normalize_name("{{2 lines of comma separated skills}}") == "2_LINES_OF_COMMA_SEPARATED_SKILLS"
    assert normalize_name("{{Methods, Systems & Technologies }}") == "METHODS_SYSTEMS_TECHNOLOGIES"


def test_scan_finds_split_runs_and_occurrences():
    doc = build_fixture_doc()
    found = {(p.name, p.occurrence) for p in scan(doc)}
    assert ("RANK", 0) in found
    assert ("RANK", 1) in found  # the one in the table
    assert ("JOB_RELEVANT_TECHNOLOGIES", 0) in found
    assert ("AREA_OF_EMPHASIS", 0) in found
    assert ("AREA_OF_EMPHASIS", 1) in found
    assert ("2_LINES_OF_COMMA_SEPARATED_SKILLS", 0) in found
    assert ("2_LINES_OF_COMMA_SEPARATED_SKILLS", 1) in found


def test_single_brace_flagged():
    doc = build_fixture_doc()
    by_name = {}
    for p in scan(doc):
        by_name.setdefault(p.name, p)
    assert by_name["AREA_OF_EMPHASIS"].single_brace is True
    assert by_name["RANK"].single_brace is False


def test_fill_replaces_everything_and_preserves_formatting():
    doc = build_fixture_doc()
    values = {
        ("RANK", 0): "Senior",
        ("RANK", 1): "Lead",
        ("JOB_RELEVANT_TECHNOLOGIES", 0): "Python, AWS",
        ("AREA_OF_EMPHASIS", 0): "Cloud",
        ("AREA_OF_EMPHASIS", 1): "Security",
        ("2_LINES_OF_COMMA_SEPARATED_SKILLS", 0): "Python, SQL",
        ("2_LINES_OF_COMMA_SEPARATED_SKILLS", 1): "Docker, Kubernetes",
    }
    fill(doc, values)

    assert scan(doc) == []
    assert doc.paragraphs[0].text == "Hello Senior engineer"
    split_para = doc.paragraphs[1]
    assert split_para.text == "Experienced with Python, AWS and more"
    assert split_para.runs[0].bold is True
    assert doc.paragraphs[2].text == "(Cloud, Security)"
    assert doc.paragraphs[3].text == "Python, SQL"
    assert doc.paragraphs[4].text == "Docker, Kubernetes"
    assert doc.tables[0].rows[0].cells[0].text == "Cell rank: Lead"


def test_fill_distinct_values_per_occurrence():
    doc = build_fixture_doc()
    fill(doc, {("RANK", 0): "First", ("RANK", 1): "Second"})
    assert doc.paragraphs[0].text == "Hello First engineer"
    assert doc.tables[0].rows[0].cells[0].text == "Cell rank: Second"
    # Unfilled placeholders remain untouched.
    assert "{{JOB_RELEVANT" in doc.paragraphs[1].text


def test_multiline_value_becomes_line_break():
    doc = docx.Document()
    doc.add_paragraph("Skills: {{2 lines of comma separated skills}}")
    fill(doc, {("2_LINES_OF_COMMA_SEPARATED_SKILLS", 0): "line one\nline two"})
    assert doc.paragraphs[0].text == "Skills: line one\nline two"
    assert scan(doc) == []
