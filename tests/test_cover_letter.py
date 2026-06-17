"""Cover-letter service + endpoint: bundled template, target field expansion,
precedence, and the bundled asset's integrity."""
import base64
import io
import json

import docx
import pytest

from app import create_app
from tailor.docxio.scanner import scan
from tailor.paths import DEFAULT_COVER_LETTER_TEMPLATE
from tailor.service import tailor_cover_letter
from tests.test_service import POSTING, make_docx


def result_text(docx_bytes: bytes) -> str:
    document = docx.Document(io.BytesIO(docx_bytes))
    return "\n".join(p.text for p in document.paragraphs)


# ----- bundled asset -----

def test_bundled_template_has_expected_placeholders():
    document = docx.Document(DEFAULT_COVER_LETTER_TEMPLATE)
    names = {p.name for p in scan(document)}
    assert {"TARGET_ROLE", "TARGET_ORGANIZATION", "FULL_NAME", "CURRENT_EMPLOYER",
            "YEARS_OF_EXPERIENCE", "JOB_RELEVANT_TECHNOLOGIES",
            "ORGANIZATION_CONTEXT", "ROLE_FOCUS"} <= names


# ----- service -----

def test_cover_letter_uses_bundled_template_by_default():
    docx_bytes, report = tailor_cover_letter(
        POSTING, target_role="Digital Content Director",
        target_organization="Franklin & Marshall College")
    text = result_text(docx_bytes)
    assert "Digital Content Director" in text
    assert "Franklin & Marshall College" in text
    assert "Alex Shaw" in text          # FULL_NAME from profile.json
    assert "Mutual of Omaha" in text    # CURRENT_EMPLOYER from profile.json


def test_target_fields_expand_to_all_occurrences():
    docx_bytes, report = tailor_cover_letter(
        POSTING, target_organization="Acme University")
    # The org appears 5x in the letter; none should survive as a placeholder.
    assert result_text(docx_bytes).count("Acme University") >= 4
    org_unfilled = [k for k in report["unfilled"] if k.startswith("TARGET_ORGANIZATION")]
    assert org_unfilled == []
    sources = {s["key"]: s["source"] for s in report["slots"]}
    assert sources["TARGET_ORGANIZATION::0"] == "field"


def test_values_override_beats_field_value():
    template = make_docx("Role: {{TARGET_ROLE}} / {{TARGET_ROLE}}")
    docx_bytes, _ = tailor_cover_letter(
        POSTING, docx_bytes=template, target_role="Director",
        values={"TARGET_ROLE::1": "VP"})
    text = result_text(docx_bytes)
    assert "Role: Director / VP" in text


def test_uploaded_template_overrides_bundled():
    template = make_docx("Hi {{TARGET_ORGANIZATION}}")
    docx_bytes, _ = tailor_cover_letter(
        POSTING, docx_bytes=template, target_organization="Globex")
    assert result_text(docx_bytes).strip() == "Hi Globex"


# ----- research-fed slots (A) -----

def test_cover_letter_field_values_compose_from_structured_facts():
    from tailor.compose import cover_letter_field_values
    fv = cover_letter_field_values({
        "company": {"name": "Acme", "industry": "Insurance"},
        "role": {"essential_skills": ["stakeholder management", "agile delivery",
                                      "cloud architecture", "extra"]}})
    assert fv["ORGANIZATION_CONTEXT"] == "your work in Insurance"
    assert fv["ROLE_FOCUS"] == "stakeholder management, agile delivery, and cloud architecture"


def test_cover_letter_field_values_partial_and_empty():
    from tailor.compose import cover_letter_field_values
    assert cover_letter_field_values({}) == {}
    assert cover_letter_field_values({"company": {"industry": "Finance"}}) == {
        "ORGANIZATION_CONTEXT": "your work in Finance"}


def test_extra_field_values_render_into_document():
    docx_bytes, _ = tailor_cover_letter(
        POSTING, target_role="Director", target_organization="Acme",
        extra_field_values={"ORGANIZATION_CONTEXT": "your work in Insurance",
                            "ROLE_FOCUS": "stakeholder management and delivery"})
    text = result_text(docx_bytes)
    assert "your work in Insurance" in text
    assert "stakeholder management and delivery" in text
    assert "{{ORGANIZATION_CONTEXT}}" not in text


# ----- endpoint -----

@pytest.fixture()
def client():
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def test_endpoint_no_template_uses_bundled(client):
    res = client.post("/api/v1/cover-letter", data={
        "posting": POSTING,
        "target_role": "Digital Content Director",
        "target_organization": "Franklin & Marshall College",
    }, content_type="multipart/form-data")
    assert res.status_code == 200
    assert res.headers["Content-Disposition"].endswith('"Cover Letter.docx"')
    text = result_text(res.data)
    assert "Digital Content Director" in text
    assert "Franklin & Marshall College" in text


def test_endpoint_json_base64_and_report(client):
    res = client.post("/api/v1/cover-letter", json={
        "posting": POSTING,
        "target_role": "Engineer",
        "target_organization": "Initech",
    }, headers={"Accept": "application/json"})
    assert res.status_code == 200
    body = res.get_json()
    assert base64.b64decode(body["docx_b64"])[:2] == b"PK"
    assert "unfilled" in body["report"]


def test_endpoint_uploaded_template(client):
    template = make_docx("Dear {{TARGET_ORGANIZATION}} team")
    res = client.post("/api/v1/cover-letter", data={
        "posting": POSTING,
        "template": (io.BytesIO(template), "cl.docx"),
        "target_organization": "Stark Industries",
    }, content_type="multipart/form-data")
    assert "Stark Industries" in result_text(res.data)


def test_endpoint_requires_posting(client):
    res = client.post("/api/v1/cover-letter", json={
        "target_role": "x", "target_organization": "y"})
    assert res.status_code == 400
    assert res.get_json()["error"] == "InvalidInputError"
