"""API-level workflow behavior: legacy default, composed opt-in, compare, health."""
import base64
import io
import json

import docx
import pytest

import tailor.web.api as api
from app import create_app
from tailor.clients.base import DownstreamError
from tests.fakes import NEWS_DATA, FakeGenerator, FakeParser, FakeResearcher
from tests.test_service import POSTING, TEMPLATE


@pytest.fixture()
def client():
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def multipart(**extra):
    data = {"posting": POSTING, "template": (io.BytesIO(TEMPLATE), "t.docx")}
    data.update(extra)
    return {"data": data, "content_type": "multipart/form-data"}


def use_fakes(monkeypatch, parser=None, researcher=None, generator=None):
    monkeypatch.setattr(api, "get_parser_client", lambda: parser)
    monkeypatch.setattr(api, "get_researcher_client", lambda: researcher)
    monkeypatch.setattr(api, "get_generator_client", lambda: generator)


def test_default_workflow_is_legacy(client, monkeypatch):
    # Even with a parser configured, the default workflow must stay legacy and
    # not call it.
    parser = FakeParser()
    use_fakes(monkeypatch, parser=parser)
    body = client.post("/api/v1/proposals", **multipart()).get_json()
    assert body["workflow"] == "legacy"
    assert parser.calls == 0


def test_composed_uses_parser_and_returns_research(client, monkeypatch):
    parser = FakeParser()
    researcher = FakeResearcher(results_by_intent={
        "concept.overview": {"summary": "context"}})
    use_fakes(monkeypatch, parser=parser, researcher=researcher)
    body = client.post("/api/v1/proposals", **multipart(workflow="composed")).get_json()
    assert body["workflow"] == "composed"
    assert parser.calls == 1
    assert body["meta"]["degraded"] is False
    assert body["research"]  # advisory suggestions present


def test_proposals_composed_returns_company_news(client, monkeypatch):
    researcher = FakeResearcher(results_by_intent={
        "concept.overview": {"summary": "ctx"}, "company.news": NEWS_DATA})
    use_fakes(monkeypatch, parser=FakeParser(), researcher=researcher)
    body = client.post("/api/v1/proposals", **multipart(
        workflow="composed", target_organization="Acme Insurance Group")).get_json()
    assert body["company_news"]["articles"]
    assert all(a["tone"] >= 2.0 for a in body["company_news"]["articles"])


def test_proposals_composed_no_company_without_target(client, monkeypatch):
    researcher = FakeResearcher(results_by_intent={"company.news": NEWS_DATA})
    use_fakes(monkeypatch, parser=FakeParser(), researcher=researcher)
    body = client.post("/api/v1/proposals", **multipart(workflow="composed")).get_json()
    assert body["company_news"] == {}
    assert researcher.research_calls == 0  # no target -> no news call


def test_composed_falls_back_when_parser_down(client, monkeypatch):
    parser = FakeParser(error=DownstreamError("parser", "boom"))
    use_fakes(monkeypatch, parser=parser, researcher=None)
    body = client.post("/api/v1/proposals", **multipart(workflow="composed")).get_json()
    assert body["workflow"] == "composed"
    assert body["meta"]["degraded"] is True
    assert body["slots"]  # still produced proposals via local fallback


def test_tailor_composed_does_not_depend_on_researcher(client, monkeypatch):
    # If research were in the resume path, a raising researcher would surface;
    # it must not — resume output is deterministic and research-free.
    class Exploding:
        def batch(self, requests):
            raise AssertionError("researcher must not be called from /tailor")

        def health(self):
            return {"version": "x"}

    use_fakes(monkeypatch, parser=FakeParser(), researcher=Exploding())
    res = client.post("/api/v1/resume", **multipart(workflow="composed"))
    assert res.status_code == 200
    assert res.data[:2] == b"PK"


def test_composed_renders_via_generator(client, monkeypatch):
    generator = FakeGenerator()
    use_fakes(monkeypatch, parser=FakeParser(), generator=generator)
    res = client.post("/api/v1/resume", headers={"Accept": "application/json"},
                      **multipart(workflow="composed"))
    assert res.status_code == 200
    assert generator.calls and generator.calls[0]["document_type"] == "docx"
    assert res.get_json()["report"]["meta"]["renderer"] == "generator"


def test_legacy_never_calls_generator(client, monkeypatch):
    generator = FakeGenerator()
    use_fakes(monkeypatch, parser=FakeParser(), generator=generator)
    res = client.post("/api/v1/resume", headers={"Accept": "application/json"},
                      **multipart(workflow="legacy"))
    assert generator.calls == []
    assert res.get_json()["report"]["meta"]["renderer"] == "local"


def test_composed_generator_outage_falls_back(client, monkeypatch):
    generator = FakeGenerator(error=DownstreamError("generator", "down"))
    use_fakes(monkeypatch, parser=FakeParser(), generator=generator)
    res = client.post("/api/v1/resume", headers={"Accept": "application/json"},
                      **multipart(workflow="composed"))
    report = res.get_json()["report"]
    assert report["meta"]["renderer"] == "local"
    assert any("generator unavailable" in w for w in report.get("warnings", []))


def test_resume_json_carries_advisory_research_but_not_in_doc(client, monkeypatch):
    researcher = FakeResearcher(results_by_intent={
        "concept.overview": {"summary": "ctx"}, "company.news": NEWS_DATA})
    use_fakes(monkeypatch, parser=FakeParser(), researcher=researcher)
    # JSON response -> report carries advisory research + company news...
    js = client.post("/api/v1/resume", headers={"Accept": "application/json"},
                     **multipart(workflow="composed",
                                 target_organization="Acme Insurance Group")).get_json()
    assert js["report"].get("research")
    assert js["report"].get("company_news", {}).get("articles")
    # ...but the document is byte-identical to a run with no researcher (research
    # never enters resume output).
    use_fakes(monkeypatch, parser=FakeParser(), researcher=None)
    plain = client.post("/api/v1/resume", headers={"Accept": "application/json"},
                        **multipart(workflow="composed",
                                    target_organization="Acme Insurance Group")).get_json()
    assert js["docx_b64"] == plain["docx_b64"]


def test_resume_binary_makes_no_researcher_call(client, monkeypatch):
    researcher = FakeResearcher(results_by_intent={"concept.overview": {"summary": "x"}})
    use_fakes(monkeypatch, parser=FakeParser(), researcher=researcher)
    res = client.post("/api/v1/resume", **multipart(workflow="composed",
                                                    target_organization="Acme"))
    assert res.status_code == 200
    assert res.data[:2] == b"PK"
    assert researcher.batch_calls == 0 and researcher.research_calls == 0


def test_unknown_workflow_is_400(client, monkeypatch):
    use_fakes(monkeypatch, parser=FakeParser())
    res = client.post("/api/v1/proposals", **multipart(workflow="banana"))
    assert res.status_code == 400


def test_compare_endpoint(client, monkeypatch):
    use_fakes(monkeypatch, parser=FakeParser())
    body = client.post("/api/v1/compare", **multipart()).get_json()
    assert "slots" in body and body["slots"]
    assert body["composed_meta"]["workflow"] == "composed"
    assert isinstance(body["changed_count"], int)


def test_cover_letter_composed_research_in_report(client, monkeypatch):
    researcher = FakeResearcher(results_by_intent={
        "company.profile": {"name": "Acme University"},
        "role.responsibilities": {"title": "Director"}})
    use_fakes(monkeypatch, parser=FakeParser(), researcher=researcher)
    res = client.post("/api/v1/cover-letter", headers={"Accept": "application/json"},
                      data={"posting": POSTING, "workflow": "composed",
                            "target_role": "Director",
                            "target_organization": "Acme University"},
                      content_type="multipart/form-data")
    report = res.get_json()["report"]
    assert report["research"]["company"]["name"] == "Acme University"
    assert report["meta"]["workflow"] == "composed"


def test_cover_letter_composed_survives_researcher_outage(client, monkeypatch):
    researcher = FakeResearcher(error=DownstreamError("researcher", "down"))
    use_fakes(monkeypatch, parser=FakeParser(), researcher=researcher)
    res = client.post("/api/v1/cover-letter",
                      data={"posting": POSTING, "workflow": "composed",
                            "target_role": "Director",
                            "target_organization": "Acme University"},
                      content_type="multipart/form-data")
    assert res.status_code == 200
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(res.data)).paragraphs)
    assert "Acme University" in text          # target field still filled
    assert "{{ORGANIZATION_CONTEXT}}" in text  # research slot stays a visible placeholder


def test_cover_letter_composed_renders_research_facts(client, monkeypatch):
    researcher = FakeResearcher(results_by_intent={
        "company.profile": {"name": "Acme University", "industry": "Higher Education"},
        "role.responsibilities": {"title": "Director",
                                  "essential_skills": ["digital strategy", "UX direction"]},
        "company.news": NEWS_DATA})
    use_fakes(monkeypatch, parser=FakeParser(), researcher=researcher)
    res = client.post("/api/v1/cover-letter", headers={"Accept": "application/json"},
                      data={"posting": POSTING, "workflow": "composed",
                            "target_role": "Director",
                            "target_organization": "Acme University"},
                      content_type="multipart/form-data")
    body = res.get_json()
    text = "\n".join(
        p.text for p in docx.Document(io.BytesIO(base64.b64decode(body["docx_b64"]))).paragraphs)
    assert "your work in Higher Education" in text
    assert "digital strategy and UX direction" in text
    assert set(body["report"]["research"]["applied_slots"]) == {
        "ORGANIZATION_CONTEXT", "ROLE_FOCUS"}
    assert body["report"]["research"]["news"]["articles"]  # news stays advisory in report


def test_health_reports_workflow_and_downstream(client, monkeypatch):
    use_fakes(monkeypatch, parser=FakeParser(), researcher=FakeResearcher())
    body = client.get("/api/v1/health").get_json()
    assert body["default_workflow"] == "legacy"
    assert body["downstream"]["parser"]["configured"] is True
    assert body["downstream"]["parser"]["version"] == "0.3.0"
