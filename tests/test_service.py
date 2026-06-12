import io

import docx
import pytest

from tailor.docxio.scanner import scan
from tailor.service import (InvalidDocxError, InvalidInputError,
                            NoPlaceholdersError, propose_for, tailor_document)

POSTING = "We need CI/CD, Kubernetes, and Python experience plus leadership."


def make_docx(*paragraphs: str) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


TEMPLATE = make_docx("Hello {{RANK}}", "Impact: {{MEASURABLE_IMPACT}}",
                     "Tech: {{JOB_RELEVANT_TECHNOLOGIES}}")


def test_propose_for_returns_slots_and_keywords():
    slots, keywords = propose_for(POSTING, TEMPLATE)
    keys = [slot.key for slot in slots]
    assert keys == ["RANK::0", "MEASURABLE_IMPACT::0", "JOB_RELEVANT_TECHNOLOGIES::0"]
    assert any(kw.canonical == "Kubernetes" for kw in keywords)


def test_tailor_uses_proposals_and_overrides():
    result_bytes, report = tailor_document(
        POSTING, TEMPLATE, values={"MEASURABLE_IMPACT::0": "a 9x speedup"})
    document = docx.Document(io.BytesIO(result_bytes))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "a 9x speedup" in text
    assert "{{MEASURABLE_IMPACT}}" not in text
    by_key = {slot["key"]: slot for slot in report["slots"]}
    assert by_key["MEASURABLE_IMPACT::0"]["source"] == "overridden"
    assert by_key["RANK::0"]["source"] == "proposed"


def test_empty_override_keeps_placeholder_and_reports_unfilled():
    result_bytes, report = tailor_document(
        POSTING, TEMPLATE, values={"RANK::0": ""})
    document = docx.Document(io.BytesIO(result_bytes))
    assert any(p.name == "RANK" for p in scan(document))
    assert "RANK::0" in report["unfilled"]


def test_request_profile_and_library_override_bundled():
    template = make_docx("{{RANK}} and {{ACTION}}")
    _, report = tailor_document(
        POSTING, template,
        profile={"RANK": "Distinguished"},
        library=[{"text": "Spearheaded", "slots": ["ACTION"], "tags": ["Leadership"]}])
    by_key = {slot["key"]: slot for slot in report["slots"]}
    assert by_key["RANK::0"]["final_value"] == "Distinguished"
    assert by_key["ACTION::0"]["final_value"] == "Spearheaded"


def test_errors():
    with pytest.raises(InvalidInputError):
        propose_for("   ", TEMPLATE)
    with pytest.raises(InvalidDocxError):
        propose_for(POSTING, b"not a docx")
    with pytest.raises(NoPlaceholdersError):
        propose_for(POSTING, make_docx("plain document, no slots"))
    with pytest.raises(InvalidInputError):
        tailor_document(POSTING, TEMPLATE, values=["not", "a", "dict"])
    with pytest.raises(InvalidInputError):
        propose_for(POSTING, TEMPLATE, library=[{"tags": ["no text"]}])


def test_deterministic_output_bytes_report():
    _, first = tailor_document(POSTING, TEMPLATE)
    _, second = tailor_document(POSTING, TEMPLATE)
    assert first == second
