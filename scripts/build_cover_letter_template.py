"""Generate data/cover_letter_template.docx from the templated cover letter.

Reproducible: run `python scripts/build_cover_letter_template.py` to regenerate.
The placeholders mirror the resume's vocabulary so the same fill engine works;
TARGET_ROLE / TARGET_ORGANIZATION are supplied per-request by the cover-letter
endpoint, and FULL_NAME / CURRENT_EMPLOYER come from profile.json.
"""
import sys
from pathlib import Path

import docx

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from tailor.paths import DEFAULT_COVER_LETTER_TEMPLATE  # noqa: E402

PARAGRAPHS = [
    "Dear Hiring Committee,",

    "With {{YEARS_OF_EXPERIENCE}}+ years of experience leading {{SCALE_DESCRIPTOR}} "
    "{{SOLUTION_TYPES}} supporting {{USER_SCALE}} internal users and {{EVENT_SCALE}} "
    "daily operational events while managing {{INITIATIVE_TYPE}} initiatives and "
    "delivering {{DOMAIN_CAPABILITIES}} instruction to 100+ students per term, I am "
    "excited to apply for the {{TARGET_ROLE}} position at {{TARGET_ORGANIZATION}}. My "
    "background combines {{SOLUTION_TYPES}}, {{TECHNICAL_CAPABILITIES}}, "
    "{{DOMAIN_CAPABILITIES}}, {{LEADERSHIP_CAPABILITIES}}, and higher education "
    "instruction in ways that align strongly with {{TARGET_ORGANIZATION}}'s vision "
    "for {{STRATEGIC_OUTCOME}}.",

    "In my current role at {{CURRENT_EMPLOYER}}, I lead {{INITIATIVE_TYPE}} and "
    "{{SOLUTION_OR_INITIATIVE}} focused on {{TECHNICAL_CAPABILITIES}}, "
    "{{DOMAIN_CAPABILITIES}}, and {{JOB_RELEVANT_SOLUTIONS}} supporting critical "
    "operational systems. I oversee projects involving {{JOB_RELEVANT_TECHNOLOGIES}} "
    "and {{TECHNICAL_CAPABILITIES}} while collaborating closely with "
    "{{SCOPE_OR_STAKEHOLDERS}} to {{ACTION_RESULT}} across enterprise platforms. A "
    "significant portion of my work involves {{ACTION_OR_IMPLEMENTATION}} that "
    "{{TECHNICAL_OR_BUSINESS_RESULT}}.",

    "Beyond {{INITIATIVE_TYPE}} leadership, I bring extensive experience developing "
    "and managing digital learning and communication environments through my work as "
    "an adjunct professor in {{AREA_OF_EMPHASIS}} and {{AREA_OF_EMPHASIS}}. I have "
    "designed and modernized project-based curriculum focused on {{LIST OF 4 COURSE "
    "TOPICS RELEVANT TO JOB POSTING - PRIORITIZE TECHNOLOGIES, PEOPLE SKILLS}}. This "
    "experience strengthened my ability to communicate technical concepts clearly, "
    "collaborate with diverse audiences, and {{STRATEGIC_OUTCOMES}}.",

    "What draws me most to {{TARGET_ORGANIZATION}} is {{ORGANIZATION_CONTEXT}} and the "
    "opportunity to combine strategic {{DOMAIN_CAPABILITIES}} leadership with meaningful "
    "institutional impact. I am especially drawn to the role's focus on {{ROLE_FOCUS}}, "
    "and excited by its emphasis on {{AREAS_OF_EMPHASIS}}, {{DELIVERY_PRACTICES}}, "
    "{{LEADERSHIP_CAPABILITIES}}, and {{SOLUTION_OR_CAPABILITY}}. My experience leading "
    "technical initiatives while balancing {{DOMAIN_CAPABILITIES}} and organizational "
    "priorities has prepared me to contribute effectively to {{TARGET_ORGANIZATION}}'s "
    "mission.",

    "I would welcome the opportunity to further discuss how my experience in "
    "{{SOLUTION_TYPES}}, {{INITIATIVE_TYPE}}, {{DOMAIN_CAPABILITIES}}, and higher "
    "education instruction can support {{TARGET_ORGANIZATION}}'s continued "
    "{{STRATEGIC_OUTCOME}}. Thank you for your time and consideration.",

    "Sincerely,",
    "{{FULL_NAME}}",
]


def build() -> None:
    document = docx.Document()
    for text in PARAGRAPHS:
        document.add_paragraph(text)
    DEFAULT_COVER_LETTER_TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
    document.save(DEFAULT_COVER_LETTER_TEMPLATE)
    print(f"wrote {DEFAULT_COVER_LETTER_TEMPLATE}")


if __name__ == "__main__":
    build()
