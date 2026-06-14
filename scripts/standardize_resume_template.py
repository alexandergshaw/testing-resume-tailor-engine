"""Rewrite a resume template's placeholders to consistent {{UPPER_SNAKE_CASE}}.

Reads any source .docx, renames every placeholder to its canonical form
(preserving all formatting via the engine's cross-run replacer), and writes
data/resume_template.docx. Idempotent: running it on an already-standardized
file is a no-op.

Usage:
    python scripts/standardize_resume_template.py [SOURCE.docx]

SOURCE defaults to the bundled data/resume_template.docx so re-runs are safe.
"""
import re
import sys
from pathlib import Path

import docx

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from tailor.docxio.filler import fill          # noqa: E402
from tailor.docxio.scanner import scan          # noqa: E402
from tailor.paths import DEFAULT_RESUME_TEMPLATE  # noqa: E402

_COURSE_RE = re.compile(r"^LIST_OF_(\d+)_COURSE_TOPICS")

# Verbose/sentence placeholders get concise names; everything else uses its
# already-normalized name (normalize_name() handles Title Case, spaces,
# single braces, etc.).
_RENAMES = {"2_LINES_OF_COMMA_SEPARATED_SKILLS": "SKILLS_LINE"}


def canonical_name(normalized: str) -> str:
    if normalized in _RENAMES:
        return _RENAMES[normalized]
    course = _COURSE_RE.match(normalized)
    if course:
        return f"COURSE_TOPICS_{course.group(1)}"
    return normalized


def standardize(source: Path, dest: Path) -> dict[str, str]:
    document = docx.Document(source)
    placeholders = scan(document)
    # Replace each placeholder token with its canonical {{NAME}} form.
    values = {(ph.name, ph.occurrence): "{{" + canonical_name(ph.name) + "}}"
              for ph in placeholders}
    fill(document, values)
    dest.parent.mkdir(parents=True, exist_ok=True)
    document.save(dest)
    # Return the raw -> canonical mapping for reporting.
    return {ph.raw: "{{" + canonical_name(ph.name) + "}}" for ph in placeholders}


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_RESUME_TEMPLATE
    mapping = standardize(src, DEFAULT_RESUME_TEMPLATE)
    print(f"standardized {src} -> {DEFAULT_RESUME_TEMPLATE}")
    for old, new in sorted(set(mapping.items())):
        if old != new:
            print(f"  {old}  ->  {new}")
